"""
PulseCare AI - Microsoft Word Project Report Generator
Modeled exactly after the LPU CSE494 MongoDB Project Report reference PDF.
Outputs: /Users/neerajnegi/Desktop/PulseCare_AI_MongoDB_Project_Report.docx
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image, ImageDraw, ImageFont

OUTPUT_PATH = "/Users/neerajnegi/Desktop/PulseCare_AI_MongoDB_Project_Report.docx"
SCREENSHOTS_DIR = "/Users/neerajnegi/Desktop/mongodb_project/reports/screenshots"

# Color Palette (derived from reference PDF)
COLOR_PRIMARY_NAVY = RGBColor(11, 27, 61)     # #0B1B3D
COLOR_SECONDARY_BLUE = RGBColor(2, 132, 199)  # #0284C7
COLOR_DARK_TEXT = RGBColor(30, 41, 59)        # #1E293B
COLOR_MUTED_GREY = RGBColor(100, 116, 139)    # #64748B
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_NAVY = "0B1B3D"
HEX_LIGHT_BLUE = "F0F9FF"
HEX_BORDER_BLUE = "0284C7"
HEX_ALT_ROW = "F8FAFC"
HEX_HEADER_BG = "0B1B3D"
HEX_BORDER_GREY = "CBD5E1"

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color=HEX_BORDER_GREY, sz="4", val="single"):
    """Sets delicate subtle borders for clean academic look."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, text, bold_prefix="Key contribution: "):
    """Creates a full-width callout box with light blue background and border, matching the PDF."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, HEX_LIGHT_BLUE)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border: left accent bar or subtle all-around
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:color="{HEX_BORDER_BLUE}"/>'
        f'<w:left w:val="single" w:sz="18" w:color="{HEX_BORDER_BLUE}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:color="{HEX_BORDER_BLUE}"/>'
        f'<w:right w:val="single" w:sz="6" w:color="{HEX_BORDER_BLUE}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    r_bold = p.add_run(bold_prefix)
    r_bold.font.name = 'Arial'
    r_bold.font.size = Pt(9.5)
    r_bold.font.bold = True
    r_bold.font.color.rgb = COLOR_PRIMARY_NAVY
    
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK_TEXT
    
    # Empty space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

def add_page_number(run):
    """Inserts a dynamic PAGE number field into a run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def generate_arch_diagram_image():
    """Generates the architecture diagram matching the cover page diagram in reference PDF."""
    img = Image.new("RGB", (1000, 360), color="#F8FAFC")
    draw = ImageDraw.Draw(img)
    
    # Container box
    draw.rounded_rectangle([(10, 10), (990, 350)], radius=12, fill="#FFFFFF", outline="#E2E8F0", width=2)
    
    # Top banner title
    draw.rectangle([(10, 10), (990, 60)], fill="#F1F5F9")
    draw.text((360, 24), "PulseCare AI - Clinical Telemetry Architecture", fill="#1E293B")
    
    # Box 1: Frontend (Blue/Cyan)
    draw.rounded_rectangle([(40, 100), (300, 200)], radius=10, fill="#0284C7")
    draw.text((100, 130), "Clinical Client UI", fill="#FFFFFF")
    draw.text((80, 155), "HTML5 • Modern Vanilla JS", fill="#E0F2FE")
    
    # Arrow 1: Frontend to Flask
    draw.line([(300, 150), (370, 150)], fill="#64748B", width=3)
    draw.polygon([(370, 145), (385, 150), (370, 155)], fill="#64748B")
    
    # Box 2: Flask API (Emerald)
    draw.rounded_rectangle([(390, 100), (650, 200)], radius=10, fill="#0D9488")
    draw.text((450, 130), "Flask REST API", fill="#FFFFFF")
    draw.text((440, 155), "Python 3.13 • PyMongo", fill="#CCFBF1")
    
    # Arrow 2: Flask to Mongo
    draw.line([(650, 150), (720, 150)], fill="#64748B", width=3)
    draw.polygon([(720, 145), (735, 150), (720, 155)], fill="#64748B")
    
    # Box 3: MongoDB (Amber/Brown)
    draw.rounded_rectangle([(740, 100), (960, 200)], radius=10, fill="#D97706")
    draw.text((790, 130), "MongoDB 7.0 Engine", fill="#FFFFFF")
    draw.text((775, 155), "$jsonSchema • 5 Pipelines", fill="#FEF3C7")
    
    # Lower Box: Security & Validation
    draw.rounded_rectangle([(400, 230), (640, 290)], radius=8, fill="#1E293B")
    draw.text((425, 245), "$jsonSchema Validation", fill="#FFFFFF")
    draw.text((440, 265), "Error 121 Rejection", fill="#94A3B8")
    draw.line([(520, 200), (520, 230)], fill="#64748B", width=2)
    
    # Footer subtitle note
    draw.text((190, 315), "High-throughput bedside telemetry ingestion; atomic $push vitals stream & $facet analytics", fill="#64748B")
    
    diag_path = "/Users/neerajnegi/Desktop/mongodb_project/reports/screenshots/architecture_diagram.png"
    img.save(diag_path, "PNG")
    return diag_path

print("Helper definitions loaded.")
